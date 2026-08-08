"""Split from the former single-file module. Re-exports every name
the flat module exposed, so `from fdd_utils.ui import X` is unchanged."""
from __future__ import annotations

# namespace parity: the pre-split flat module re-exported everything it
# imported at module level. Reproduced so `from fdd_utils.ui import X`
# resolves for X that were only ever incidental re-exports.
from typing import Any
from datetime import datetime
from typing import Any, Callable, Dict, Iterable, List, Optional
import pandas as pd
from ..financial_common import (
    build_income_statement_period_label,
    clean_english_placeholders,
    dedupe_non_empty,
    visible_descriptions,
)
from ..workbook import (
    INTERNAL_ROW_KEY,
    split_accounts_by_type as shared_split_accounts_by_type,
    split_bilingual_entity_name,
)
import time
import traceback
from typing import Any, Dict
import streamlit as st
from ..ai import (
    FDDConfig,
    WORKBENCH_AVAILABLE_MODELS,
    build_highlighted_commentary_html,
    get_default_config_path,
    get_prompt_engine,
    is_provider_ready,
    load_yaml_config,
    parse_validator_response,
    run_ai_pipeline_with_progress,
    run_generator_reprompt,
)
from ..financial_common import extract_result_text_content, get_pipeline_result_text
from ..financial_display_format import prepare_display_dataframe, stringify_display_dataframe
from ..workbook import (
    build_account_mapping_diagnostics,
    build_workbook_preflight,
    extract_entity_names_from_preflight,
    find_mapping_key,
    get_effective_mappings,
    get_financial_sheet_options,
    load_mappings,
    split_accounts_by_type,
)
from typing import Any, List
from ..workbook import find_mapping_key, load_mappings
from datetime import timedelta
import hashlib
from pathlib import Path
import re
import tempfile
from typing import Any, Callable
from ..workbook import clear_table_inspection_cache, clear_workbook_caches
import datetime as dt_module
import logging
import os
from typing import Any, Callable, Dict, List, Optional
from ..pptx import build_pptx_structured_payloads
from ..workbook import find_mapping_key, get_effective_mappings, load_mappings
import io as _bridge_lab_io
from openpyxl import Workbook as _bridge_lab_Workbook
from openpyxl import load_workbook as _bridge_lab_load_workbook
from ..bridge_chart_prototype import build_excel_waterfall_chart, find_bridge_blocks
from ..generate_bridge_waterfall_batch import build_bridges_for_ab_tab  # noqa: F401


from .state import (  # noqa: F401
    DEFAULT_SESSION_STATE,
    DELETE_SESSION_KEYS,
    RESET_SESSION_KEYS,
    init_session_state,
    initialize_app_state,
    reset_processing_session_state,
)
from .views import (  # noqa: F401
    build_account_display_dataframe,
    build_entity_selector_model,
    build_processed_display_groups,
    derive_reconciliation_matched_keys,
    describe_statement_period,
    detect_statement_mode,
    should_render_preprocess_controls,
    _build_rhs_display_dataframe,
    _normalize_display_label,
)
from .ai_panel import (  # noqa: F401
    build_selected_pipeline_dfs,
    effective_mappings_from_session,
    extract_account_remarks_context,
    extract_result_text,
    extract_validator_metadata,
    format_dataframe_for_display,
    get_account_dataframe,
    get_financial_account_options,
    has_meaningful_result_text,
    hydrate_nested_agent_outputs,
    render_account_remarks_context,
    render_ai_generation_section,
    render_generated_content,
    _EMPTY_RESULT_MARKERS,
    _PROMPT_MANAGER,
    _result_has_pipeline_content,
    _run_demo_ai,
)
from .processed import (  # noqa: F401
    filter_reconciliation_display_rows,
    reconciliation_warning_row_count,
    render_data_tables_section,
    render_processed_view,
    render_reconciliation_metrics,
    render_reconciliation_section,
    _RECON_DISPLAY_COLUMN_MAP,
    _render_resolver_diagnostics,
    _render_single_reconciliation_tab,
    _trim_reconciliation_columns_for_display,
)
from .sidebar import (  # noqa: F401
    cleanup_stale_uploads,
    persist_uploaded_workbook,
    render_language_selector,
    render_sidebar_upload,
    _build_model_choices,
    _safe_stem,
)
from .pptx_export import (  # noqa: F401
    batch_extract_entity_data,
    batch_process_entity,
    batch_run_ai_for_entity,
    generate_pptx_presentation,
    logger,
    render_bridge_lab,
    render_bridge_lab_toggle,
    _bridge_lab_render_preview_chart,
    _bridge_lab_show_block,
)
