"""Split from the former single-file module. Re-exports every name
the flat module exposed, so `from fdd_utils.workbook import X` is unchanged."""
from __future__ import annotations

# namespace parity: the pre-split flat module re-exported everything it
# imported at module level. Reproduced so `from fdd_utils.workbook import X`
# resolves for X that were only ever incidental re-exports.
from typing import Any, Dict, Iterable, List, Optional
import pandas as pd
from ..financial_common import load_yaml_file, package_file_path
from ..keyword_registry import (
    BS_END_KEYWORDS,
    BS_HEADER_KEYWORDS,
    INDICATIVE_KEYWORDS,
    IS_END_KEYWORDS,
    IS_HEADER_KEYWORDS,
    REMARK_KEYWORDS,
    SUBTOTAL_KEYWORDS,
    SUMMARY_ACCOUNT_SKIP_KEYWORDS,
    TABLE_END_KEYWORDS,
    contains_thousand_unit_marker,
)
from typing import Any, Dict, List, Optional
from functools import lru_cache
import logging
import re
import time
from typing import Any, Callable, Dict, List, Optional
from ..financial_common import cell_text
import difflib
import os
from typing import Any, Dict, Iterable, List, Optional, Sequence
from openpyxl import load_workbook
from typing import Dict, List, Tuple, Optional, Any
from dataclasses import dataclass, field
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, List, Sequence, Tuple
from typing import Any, Dict, Tuple, Optional, List
import warnings
from ..financial_common import cell_text, coerce_numeric, normalize_financial_date_label
from collections import Counter
import math
from ..financial_common import cell_text, coerce_numeric
from ..keyword_registry import UNIT_THOUSAND_MARKERS
from difflib import SequenceMatcher
import json
from typing import Any, Callable, Dict, Iterable, List, Optional, Tuple
from ..financial_common import clean_english_placeholders, load_yaml_file, normalize_financial_date_label, package_file_path
from ..financial_display_format import add_language_display_columns
from typing import Dict, Tuple, Optional
import contextlib
import io
from typing import Any, Dict  # noqa: F401


from .mapping import (  # noqa: F401
    build_account_mapping_diagnostics,
    find_mapping_key,
    get_effective_mappings,
    iter_account_mappings,
    load_mappings,
    normalize_mapping_label,
    should_skip_account_label,
    split_accounts_by_type,
)
from .analysis import (  # noqa: F401
    MATERIAL_PCT_CHANGE,
    build_significant_movements,
    build_trend_summary,
    _TOTAL_KEYWORDS,
    _change_direction,
    _period_analysis_columns,
    _select_trend_focus_row,
    _trend_direction,
)
from .inspector import (  # noqa: F401
    CANONICAL_STAGE_LABELS,
    canonical_stage_label,
    contains_indicative_text,
    contains_unit_marker,
    date_row_index,
    load_workbook_frames,
    logger,
    primary_stage_row_index,
    profile_sheet,
    profile_workbook,
    stage_row_indices,
    _TEMPLATE_NAV_EXACT_NAMES,
    _TEMPLATE_NAV_PREFIXES,
    _TEMPLATE_NAV_SUFFIXES,
    _TITLE_SKIP_VALUES,
    _UNIT_PATTERNS,
    _cell_text,
    _coerce_numeric,
    _collect_row_labels,
    _date_labels,
    _date_row_index,
    _entity_scope,
    _header_signature,
    _is_template_nav_sheet,
    _looks_like_entity_heading,
    _normalize_spaces,
    _parse_date_label,
    _sheet_kind,
    _sheet_title,
    _stage_block_heading,
    _stage_labels,
    _stage_row_index,
    _title_row_index,
    _unit_markers,
)
from .preflight import (  # noqa: F401
    build_workbook_preflight,
    extract_entity_names_from_preflight,
    get_financial_sheet_options,
    logger,
    split_bilingual_entity_name,
    suggest_rollup_sheet_for_entity,
    _CJK_RUN_RE,
    _build_workbook_preflight_cached,
    _is_empty_value,
    _is_likely_entity_name,
    _looks_like_financial_schedule_preview,
    _looks_like_generic_entity_prefix,
    _looks_like_schedule_title_prefix,
    _normalize_for_sheet_match,
    _rows_are_blank,
    _serialize_rows,
    _strip_leading_date_fragment,
    _visible_non_blank_sheets,
)
from .table_debug import (  # noqa: F401
    RowClassification,
    TableInspection,
    TableSection,
    clear_table_inspection_cache,
    clear_workbook_caches,
    get_table_inspection,
    inspect_sheet,
    inspect_workbook,
    _find_header_rows,
    _is_subtotal_or_total,
)
from .text_export import (  # noqa: F401
    TrimmedSheet,
    export_selected_tabs_to_file,
    render_selected_tabs_text,
    _column_letter,
    _is_blank_or_na_value,
    _normalize_cell_value,
    _render_sheet_table,
    _trim_sheet,
    _validate_selected_tabs,
)
from .statements import (  # noqa: F401
    extract_balance_sheet_and_income_statement,
    extract_financial_table,
    get_valid_financial_columns,
    logger,
    parse_date,
    synthesize_balance_sheet_and_income_statement,
    _POST_IS_SECTION_MARKERS,
    _RATIO_ROW_MARKERS,
    _SYNTHETIC_BS_CATEGORY_ORDER,
    _SYNTHETIC_BS_GRAND_TOTAL_GROUPS,
    _SYNTHETIC_IS_CATEGORY_ORDER,
    _build_financial_result,
    _column_numeric_profile,
    _contains_indicative_marker,
    _dedupe_date_columns,
    _extract_income_statement_directly,
    _find_best_columns_from_header,
    _find_description_column,
    _find_extended_relaxed_date_columns,
    _find_relaxed_date_columns,
    _find_section_end_row,
    _find_table_end_row,
    _get_valid_financial_columns_for_rows,
    _looks_like_remark_text,
    _scan_relaxed_date_columns,
    _select_indicative_cluster,
    _synthesize_statement,
    _synthetic_account_total_row,
    _table_end_keywords,
)
from .schedules import (  # noqa: F401
    INTERNAL_ROW_KEY,
    PREFERRED_STAGE,
    extract_presentation_detail_table,
    infer_partial_year_annualization,
    normalize_financial_schedule,
    _CARRYING_AMOUNT_LABELS,
    _GL_CODE_RE,
    _MAX_NEST_WINDOW,
    _NON_COMPOSITION_ROW_MARKERS,
    _REPORT_TITLE_RE,
    _SUBTOTAL_KEYWORDS,
    _TOTAL_KEYWORDS,
    _UNIT_MARKERS,
    _WORKING_REMARK_KEYWORDS,
    _annualization_factor,
    _auxiliary_header_context,
    _block_title_for_stage_row,
    _build_indent_signal_index,
    _build_prompt_analysis_df,
    _build_table_linked_remarks,
    _build_working_remark_note,
    _choose_projection,
    _dedupe_columns_by_key,
    _detect_implicit_breakdowns_from_sum,
    _extract_adjacent_detail_columns,
    _extract_auxiliary_check_totals,
    _extract_entity_name_from_block_title,
    _extract_indent_signal_rows,
    _extract_supporting_notes,
    _fallback_description,
    _forward_fill_stage_row,
    _infer_indent_hierarchy,
    _is_composition_label,
    _is_numeric_enough,
    _is_pure_working_artifact,
    _is_strict_entity_title_match,
    _local_date_row_index,
    _looks_like_supporting_note,
    _looks_like_working_remark,
    _multiply_factor,
    _nest_component_rows,
    _parse_statement_date_label,
    _reclassify_indent_rollup_children,
    _rollforward_header_row_index,
    _row_type,
    _select_entity_block,
    _stage_row_indices,
    _standardize_rollforward_schedule_df,
    _trim_block_end_row,
)
from .resolver import (  # noqa: F401
    resolve_ambiguous_candidate,
    resolve_workbook_mappings,
    should_use_ai_for_candidates,
    _LOW_CONFIDENCE_MATCH_FLOOR,
    _available_candidates,
    _build_candidate_map,
    _build_dynamic_mapping_config,
    _build_financial_reference_context,
    _build_sheet_candidate_for_account,
    _candidate_passes_dynamic_confirmation,
    _candidate_sheets_for_mapping,
    _candidate_strings,
    _candidate_total_values,
    _default_ai_decider,
    _discover_dynamic_sheet_resolutions,
    _extract_sheet_names_from_ai_response,
    _infer_accounting_category,
    _is_compact_cjk_label,
    _is_exact_alias_match,
    _is_summary_account_candidate,
    _iter_financial_reference_rows,
    _normalize_label,
    _normalized_total_values,
    _pick_financial_summary_sheet,
    _rank_candidates_with_financial_signals,
    _rank_mapping_candidates,
    _resolve_dynamic_sheet_mapping,
    _resolve_manual_override_target,
    _resolve_top_ranked_candidate,
    _score_candidate,
    _semantic_alignment_adjustment,
    _sheet_type_bonus,
    _should_accept_hybrid_top_candidate,
    _statement_values_for_mapping,
    _summary_row_exact_matches_sheet,
    _token_set,
)
from .databook import (  # noqa: F401
    build_dataframe_variants_from_normalized_results,
    build_dataframes_from_normalized_results,
    detect_databook_language,
    extract_data_from_excel,
    extract_normalized_data_from_excel,
    filter_detail_accounts,
    filter_zero_value_rows,
    load_mapping,
    logger,
    _detect_report_language_from_profiles,
    _update_display_description_map,
)
from .reconcile import (  # noqa: F401
    find_account_in_dfs,
    find_reconciliation_example,
    get_total_from_dfs,
    print_reconciliation_report,
    reconcile_financial_statements,
    should_skip_mapping,
    _integrity_metadata,
    _latest_column_is_partial_period,
    _normalize_account_name,
    _reconcile_against_prior_period,
    _resolve_mapping_alias,
    _should_compare_income_statement_as_absolute,
)
from .flow import (  # noqa: F401
    logger,
    process_workbook_data,
)
