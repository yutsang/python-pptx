"""Hold a generated deck against the team's own final report, account by account.

The final report is the target the tool is calibrated to: what the team kept,
added and rewrote is the clearest statement there is of what a finished
commentary looks like. This prints, for every account the two decks share:

  * the final text and the generated text one under the other;
  * the figures the final quotes that the generated text does not (amounts in
    万元/亿元/元 and percentages, compared after unit conversion, within
    rounding) -- what the team thought worth saying that the tool left out;
  * the figures the generated text quotes that the final does not -- what the
    tool said that the team did not keep;
  * a coverage ratio over every shared account, so two generated decks can be
    compared on one number before anyone reads the text.

    python ad-hoc/workbench/compare_to_final.py <generated.pptx> <final.pptx>

Free: it reads two files and calls nothing. Its output quotes client
figures and names; do not paste it anywhere public.

Accounts are found by the "■ <account> - " bullet head the generator writes,
and in the final deck by the same head or by a paragraph that opens with an
account name the generated deck uses. An account the final writes under a
different heading is reported as unmatched, not guessed at.
"""
from __future__ import annotations

import re
import sys
from typing import Dict, List, Set, Tuple

from pptx import Presentation

# 「■ 货币资金 - 」, 「■ Other receivables (cont'd) - 」, 「货币资金：」
_HEAD = re.compile(
    r"^\s*[■▪•●]\s*(.{2,40}?)\s*(?:\(cont'd\)|\(续\)|（续）)?\s+[-–—]\s+"
    r"|^\s*[■▪•●]?\s*([^\s，,。：:]{2,14})\s*(?:\(续\)|（续）)?\s*[：:]\s*"
)
_SENTENCE = re.compile(r"[。．.，,；;]")
_AMOUNT = re.compile(r"(负|-)?\s*(\d[\d,]*(?:\.\d+)?)\s*(亿元|万元|千元|元|%)")
_SCALE = {"亿元": 1e8, "万元": 1e4, "千元": 1e3, "元": 1.0}


# 「资料来源：管理层信息；审计报告」 and the bare 「管理层信息；审计报告」 a
# footnote box holds. Read as text they opened an account called 资料来源 and
# swallowed the 税金及附加 and 财务费用 bullets that followed on the first run.
_SOURCE_LINE = re.compile(r"^\s*(?:资料来源|来源|Source)\s*[：:]|^[^\d，,。]{2,12}(?:[；;][^\d，,。]{2,12})+$")


def _paragraphs(path: str) -> List[str]:
    out: List[str] = []
    for slide in Presentation(path).slides:
        # Left column before right, top before bottom. In XML order the right
        # column came first, so an account split across the two columns read
        # continuation-first: 「并于2025年12月31日进一步降至…」 opened 应付账款.
        for shape in sorted(slide.shapes, key=lambda sh: ((sh.left or 0) // 914400 >= 5, sh.top or 0)):
            # The executive-summary band speaks for the whole statement; read as
            # commentary it was appended to whichever account came before it.
            if "summary" in (shape.name or "").lower():
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text and not _SOURCE_LINE.search(text):
                        out.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                continue  # statement and detail tables: figures the reader already has
    return out


def _by_account(paragraphs: List[str], known: Set[str]) -> Dict[str, str]:
    sections: Dict[str, List[str]] = {}
    current = None
    for text in paragraphs:
        head = _HEAD.match(text)
        name = (head.group(1) or head.group(2)).strip() if head else None
        if not head and len(text) < 30 and not _SENTENCE.search(text):
            continue  # a section label or slide title, not commentary
        if name is None:
            opening = next((k for k in sorted(known, key=len, reverse=True) if text.startswith(k)), None)
            name = opening
        if name and (not known or name in known or head):
            current = name
            body = text[head.end():] if head else text
            sections.setdefault(current, []).append(body)
        elif current:
            sections[current].append(text)
    return {k: "".join(v) for k, v in sections.items()}


_AMOUNT_ENG = re.compile(r"(-)?\s*(?:CNY|RMB)\s*(-)?\s*(\d[\d,]*(?:\.\d+)?)\s*(million|mn|K|thousand)?", re.IGNORECASE)
_SCALE_ENG = {"million": 1e6, "mn": 1e6, "k": 1e3, "thousand": 1e3, "": 1.0}


def _figures(text: str) -> List[Tuple[str, float, str]]:
    found = []
    for sign_a, sign_b, number, unit in _AMOUNT_ENG.findall(text):
        value = float(number.replace(",", "")) * _SCALE_ENG[(unit or "").lower()] * (-1 if (sign_a or sign_b) else 1)
        found.append(("amt", value, f"CNY{number}{unit}"))
    for sign, number, unit in _AMOUNT.findall(text):
        value = float(number.replace(",", "")) * (-1 if sign else 1)
        if unit == "%":
            found.append(("pct", value, f"{number}%"))
        else:
            found.append(("amt", value * _SCALE[unit], f"{'-' if sign else ''}{number}{unit}"))
    return found


def _has(pool: List[Tuple[str, float, str]], kind: str, value: float) -> bool:
    for k, v, _ in pool:
        if k != kind:
            continue
        if kind == "pct" and abs(abs(v) - abs(value)) <= 0.6:
            return True
        # Rounded differently on each side (1.23亿元 against 12,300.8万元).
        if kind == "amt" and abs(value) and abs(abs(v) - abs(value)) <= max(0.006 * abs(value), 5_000):
            return True
    return False


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__)
        return 1
    ours_paragraphs, final_paragraphs = _paragraphs(sys.argv[1]), _paragraphs(sys.argv[2])
    ours = _by_account(ours_paragraphs, set())
    final = _by_account(final_paragraphs, set(ours))
    shared = [k for k in ours if k in final]
    print(f"generated: {len(ours)} account(s); final: {len(final)} section(s); shared: {len(shared)}")
    unmatched = [k for k in final if k not in ours]
    if unmatched:
        print("in the final only (heading not matched to a generated account): " + "、".join(unmatched[:20]))
    total_final = total_hit = 0
    rows = []
    for key in shared:
        f_figs, o_figs = _figures(final[key]), _figures(ours[key])
        missing = [label for kind, value, label in f_figs if not _has(o_figs, kind, value)]
        extra = [label for kind, value, label in o_figs if not _has(f_figs, kind, value)]
        hit = len(f_figs) - len(missing)
        total_final += len(f_figs)
        total_hit += hit
        rows.append((key, f_figs, hit, missing, extra))
    print(f"\nCOVERAGE: the generated deck quotes {total_hit} of the {total_final} figures the final quotes "
          f"({(100 * total_hit / total_final) if total_final else 0:.0f}%), over {len(shared)} shared account(s)\n")
    print("%-14s %8s %9s %8s" % ("account", "final#", "covered", "ours#"))
    for key, f_figs, hit, missing, extra in rows:
        print("%-14s %8d %9d %8d" % (key[:14], len(f_figs), hit, len(_figures(ours[key]))))
    for key, f_figs, hit, missing, extra in rows:
        print(f"\n{'=' * 90}\n  {key}\n{'=' * 90}")
        print(f"--- FINAL ({len(final[key])} chars) ---\n{final[key]}")
        print(f"--- GENERATED ({len(ours[key])} chars) ---\n{ours[key]}")
        if missing:
            print("    in the final, not in ours: " + "、".join(missing[:15]))
        if extra:
            print("    in ours, not in the final: " + "、".join(extra[:15]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
